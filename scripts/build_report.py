"""Regenerate the cumulative PAS report from the supplied title-page example.

Add future practical work to CONTENT. The table of contents is paginated
automatically through a preliminary render; no old page numbers are reused.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader


DEFAULT_TEMPLATE = Path("/Users/timurkamalov/Downloads/ПР1_КудзиевШД (1).docx")
DEFAULT_OUTPUT = Path(__file__).resolve().parents[2] / "Отчет_ПАС_Камалов_практики_1_2.docx"
DEFAULT_RENDERER = Path(
    "/Users/timurkamalov/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.904.11930/skills/documents/render_docx.py"
)

# The list is the content source for this report. Later practical work is added here.
CONTENT = [
    ("h1", "ВВЕДЕНИЕ"),
    ("p", "Работа относится к предметной области № 21 «Рынок цифровых активов» и посвящена спотовой паре BTC/USDT на площадке OKX. Пользователь системы должен видеть ожидаемый объём торгов и нестабильность цены на следующие календарные сутки по UTC. Прогноз предназначен для оценки необходимости повышенного внимания к инструменту, а не для выдачи рекомендации купить или продать его [1]."),
    ("p", "Выбор одной площадки делает объект измеримым: объём OKX можно получить и проверить по её собственным данным. Он не равен обороту всего мирового рынка биткоина. Вторая независимая информационная линия — официальный курс USD/RUB Банка России. Курс рассматривается как возможный внешний фактор; улучшит ли он прогноз, покажет только последующая проверка модели."),
    ("p", "Первая практическая работа определяет пользователя, показатели, горизонт, источники и ограничения. Вторая реализует получение и сохранение исходных данных. Эти работы образуют начало одного отчёта, который будет дополняться обработкой, моделями и дашбордами. На текущем этапе результатом является проверенная загрузка, а не готовый прогноз."),
    ("p", "При выборе инструмента учитывалась пригодность истории для проверки прогноза во времени. Ранее рассмотренный новый фьючерс давал лишь 212 различных торговых дней, несмотря на большее число строк с датами исполнения. Спотовая пара на OKX имеет непрерывный суточный ряд с 2018 года. Это позволит отделить ранние даты для обучения от поздних дат для проверки и сравнить модель с простым базовым прогнозом. Однако длинный ряд не гарантирует точность: рыночный режим меняется, а ошибка будущей модели пока не измерена. Наблюдаемые данные и непроверенные предположения поэтому разделены в отчёте."),

    ("h1", "1 ПРАКТИЧЕСКАЯ РАБОТА 1 ПОСТАНОВКА ЗАДАЧИ"),
    ("h2", "1.1 Объект и задача пользователя"),
    ("p", "Объект наблюдения — спотовые сделки по паре BTC/USDT на OKX. При спотовой торговле исследуется один инструмент без смены срока исполнения, характерной для фьючерсов. Начинающий аналитик или частный инвестор сможет сравнить ожидаемую активность и нестабильность с обычным уровнем прошлых дней и решить, нужно ли внимательнее наблюдать за рынком."),
    ("p", "Практическая задача ограничена предварительной оценкой риска и активности. Система не предсказывает направление цены, не гарантирует доходность и не подменяет инвестиционное решение пользователя. Поэтому качество прогноза будет оцениваться по ошибкам на будущих относительно обучения датах, а не по красивому виду графика."),

    ("h2", "1.2 Прогнозируемые показатели"),
    ("p", "После окончания суток d по UTC формируется прогноз на сутки d+1. В Таблице 1.1 приведены определения двух целевых показателей. Для обучения и проверки не допускается использовать сведения дня d+1 в его признаках."),
    ("table", ("Таблица 1.1 — Прогнозируемые показатели",
               ["Показатель", "Определение", "Единица"],
               [
                   ("Объём торгов", "Поле vol завершённой суточной свечи BTC/USDT на OKX.", "BTC"),
                   ("Дневная волатильность", "Оценка Паркинсона по логарифму отношения максимальной и минимальной цены за сутки.", "Безразмерная величина"),
               ], [4.2, 9.1, 3.7])),
    ("p", "Оценка волатильности рассчитывается после завершения дня по его максимуму и минимуму: модуль натурального логарифма отношения high к low делится на удвоенный корень из натурального логарифма двух. Это оценка ценовых колебаний внутри одних суток, а не скользящий показатель за десять дней. Такой выбор не создаёт искусственно похожие соседние целевые значения за счёт перекрывающихся окон."),
    ("p", "Для объёма используется поле vol в BTC; оборот в USDT сохраняется отдельно и не подменяет объём в базовой валюте. Прогноз на следующие сутки можно сформировать только после получения завершённой свечи текущего дня."),

    ("h2", "1.3 Источники и проверенная история"),
    ("p", "Методические указания требуют как минимум два разнородных источника с временной составляющей и достаточной историей [1]. Выбраны публичный JSON API OKX и XML-сервис Банка России. Поля и роли источников приведены в Таблице 1.2 [3, 4]."),
    ("table", ("Таблица 1.2 — Источники данных",
               ["Источник", "Данные", "Роль в системе"],
               [
                   ("OKX, BTC/USDT", "Начало суток UTC, цены открытия, максимума, минимума и закрытия; объём в BTC, оборот в USDT, признак завершённости.", "Основной ряд и оба прогнозируемых показателя."),
                   ("Банк России", "Дата действия, номинал и официальный курс USD/RUB.", "Внешний признак и контекст; его пользу нужно проверить."),
               ], [3.8, 8.6, 4.6])),
    ("p", "Публичный маршрут OKX возвращает девять элементов каждой свечи: ts, o, h, l, c, vol, volCcy, volCcyQuote и confirm. Параметр 1Dutc задаёт границу суток по UTC; значение confirm, равное единице, обозначает завершённую свечу [3]. Банк России предоставляет историю USD/RUB по коду R01235 в XML [4]."),
    ("p", "При проверке 19 сентября 2026 года у OKX обнаружены 3 174 даты с 11 января 2018 года по 19 сентября 2026 года без календарных пропусков в этом интервале. Последняя свеча ещё не была завершена. Следовательно, на момент проверки доступны 3 173 завершённых дня по 18 сентября включительно. Площадка предупреждает, что самая ранняя свеча может быть частичной; на следующем этапе она пройдёт отдельную проверку качества [3]."),

    ("h2", "1.4 Архитектура и ограничения"),
    ("p", "Архитектура состоит из трёх последовательно заполняемых слоёв. В raw лежат полученные записи и исходные JSON/XML-фрагменты. В clean данные будут типизированы, проверены и сопоставлены по датам. В mart появится одна строка на сутки UTC с признаками и целями. Журнал load_log уже фиксирует результат каждого запуска; контроль качества, прогнозы и дашборды относятся к следующим этапам [1, 2]."),
    ("p", "Крипторынок работает ежедневно, а официальный курс ЦБ не появляется на каждую календарную дату. В будущей витрине курс будет переноситься только вперёд от последней уже известной даты, с указанием возраста записи. Курс, датированный прогнозируемым днём, не должен незаметно попасть в признаки. Кроме того, USDT нельзя автоматически считать равным официальному доллару США: курс USD/RUB не применяется здесь для безусловного пересчёта цены BTC/USDT в рубли."),
    ("p", "История длиннее, чем у первоначально выбранного фьючерса, но её размер не гарантирует малую ошибку модели. Объём характеризует лишь OKX; поведение участников меняется, а курс USD/RUB может не улучшить прогноз. Тема и использование зарубежного источника согласованы студентом как план проекта; согласование с преподавателем остаётся отдельным необходимым действием."),
    ("h1", "2 ПРАКТИЧЕСКАЯ РАБОТА 2 ЗАГРУЗКА ДАННЫХ"),
    ("h2", "2.1 Состав загрузчика и хранение"),
    ("p", "Реализованы два загрузчика: дневных свечей OKX и официального курса Банка России. Они начинают запрос со следующей после последней сохранённой даты. Для OKX история читается страницами назад, а первая полная загрузка курса ЦБ разбивается на ограниченные годовые интервалы. Перед сохранением проверяется формат ответа; незавершённая свеча OKX с confirm, равным нулю, пропускается и будет запрошена позже."),
    ("p", "Таблица 2.1 показывает созданные таблицы и ключи. Вместе с выделенными полями хранится исходная строка JSON или XML, чтобы можно было восстановить происхождение значения."),
    ("table", ("Таблица 2.1 — Таблицы второй практической работы",
               ["Таблица", "Уникальный ключ", "Назначение"],
               [
                   ("raw.okx_btc_usdt_daily", "Инструмент и дата UTC", "Завершённая свеча, цены, объём, оборот и исходный JSON."),
                   ("raw.cbr_usd_rub", "Дата курса", "Курс USD/RUB, номинал и исходный XML."),
                   ("load_log", "Номер запуска", "Источник, время, статус, полученные и добавленные строки, последняя дата, ошибка."),
               ], [5.0, 4.1, 7.9])),
    ("p", "Вставка с конфликтом уникального ключа не переписывает старую строку. Временные HTTP-ошибки и превышение лимита запросов автоматически повторяются. Если загрузка всё же прерывается, текущая транзакция откатывается, а сообщение о сбое сохраняется в load_log. Это защищает raw от частично записанного запуска."),

    ("h2", "2.2 Контрольный запуск и повтор"),
    ("p", "19 сентября 2026 года загрузчики были проверены на реальных ответах обоих источников по 18 сентября включительно. Для этого использована временная SQLite-база с теми же уникальными ключами; результаты первого и повторного запуска приведены в Таблице 2.2."),
    ("table", ("Таблица 2.2 — Результат контрольной загрузки",
               ["Источник", "Первый запуск", "Повторный запуск", "Строк в БД"],
               [
                   ("OKX", "Получено и добавлено 3 173", "Получено 0, добавлено 0", "3 173"),
                   ("Банк России", "Получено и добавлено 2 147", "Получено 0, добавлено 0", "2 147"),
               ], [3.1, 5.3, 5.2, 3.4])),
    ("p", "Четыре запуска создали четыре записи в load_log. Шесть автоматических тестов прошли: они проверяют разбор XML, постраничный запрос OKX, отсутствие дублей, повторный запрос незавершённой свечи и фиксацию сбоя без частичной записи. Нулевое число добавленных строк при повторе подтверждает идемпотентность для проверенного сценария."),

    ("h2", "2.3 Границы выполненной проверки"),
    ("p", "Код рассчитан на PostgreSQL и запускается через Docker Compose после задания пароля в файле .env. На компьютере подготовки отчёта Docker отсутствует, поэтому работа контейнеров с PostgreSQL ещё не подтверждена. Проверены логика загрузчика, тесты и реальное чтение API во временной SQLite-базе; баллы за непроверенный контейнерный запуск заранее не заявляются."),
    ("p", "При контрольном прогоне через учебный прокси проверка TLS отключалась только в тестовой сессии; основной загрузчик проверяет сертификат. Далее предстоят контроль качества, слои clean и mart и сравнение прогнозных моделей по времени."),

    ("h1", "ЗАКЛЮЧЕНИЕ"),
    ("p", "В первой работе определены инструмент BTC/USDT на OKX, пользователь, два показателя и суточный горизонт. Проверка подтвердила длинную историю завершённых свечей и доступность разнородного XML-источника. Во второй работе реализована инкрементальная загрузка в raw и журнал запусков. Повторный контрольный прогон не создал дублей."),
    ("p", "Система пока не формирует прогноз: это следующий проверяемый результат после обработки данных. Длинная история устраняет главный недостаток прежней темы с новым фьючерсом, однако точность будущей модели должна быть установлена экспериментально. Отдельно остаются согласование зарубежного источника с преподавателем и проверка PostgreSQL в Docker."),

    ("h1", "СПИСОК ИСТОЧНИКОВ"),
    ("source", "1. Методические указания к проектной работе по дисциплине «Прогнозно-аналитические системы» для профиля «Управление данными». Учебный документ, предоставленный студентом."),
    ("source", "2. Критерии оценивания проектных работ по дисциплине «Прогнозно-аналитические системы». Учебный документ, предоставленный студентом."),
    ("source", "3. OKX. API guide. Candlesticks history. URL: https://app.okx.com/docs-v5/en/ (дата обращения: 19.09.2026)."),
    ("source", "4. Банк России. Получение данных с использованием XML. URL: https://www.cbr.ru/development/SXML/ (дата обращения: 19.09.2026)."),
]


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = "Times New Roman"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style(doc, name, size, *, bold=False, italic=False, first=0,
              before=0, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Times New Roman"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = italic
    s.font.color.rgb = RGBColor(0, 0, 0)
    fmt = s.paragraph_format
    fmt.alignment = align
    fmt.first_line_indent = Cm(first)
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = True
    if bold:
        fmt.keep_with_next = True


def add_paragraph(doc, value, kind="Normal"):
    p = doc.add_paragraph(value, style=kind)
    previous = p._p.getprevious()
    if previous is not None and previous.tag == qn("w:tbl"):
        p.paragraph_format.space_before = Pt(17)
    return p


def add_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")


def add_table(doc, spec):
    caption, heads, rows, widths = spec
    p = add_paragraph(doc, caption, "PAS Caption")
    p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(heads))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    add_borders(table)
    for values in [heads, *rows]:
        cells = table.rows[0].cells if values is heads else table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cells[i].width = Cm(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cells[i].paragraphs:
                para.style = doc.styles["PAS Table"]
                for run in para.runs:
                    set_font(run, 12, bold=values is heads)
    row = table.rows[0]
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def add_page_number(section):
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = False
    footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    set_font(run, 11)
    for kind, content in [("begin", None), ("instr", " PAGE "), ("separate", None),
                          ("value", "3"), ("end", None)]:
        if kind in {"begin", "separate", "end"}:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = content
        else:
            el = OxmlElement("w:t")
            el.text = content
        run._r.append(el)
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        section._sectPr.append(pg)
    pg.set(qn("w:start"), "3")


def render_report(template: Path, output: Path, group: str, teacher: str,
                  pages: dict[str, int]):
    doc = Document(template)
    body = doc._element.body
    for child in list(body)[16:-1]:
        body.remove(child)
    if "Title" not in doc.styles:
        doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    title = doc.paragraphs[4]
    title.clear()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("ОТЧЁТ ПО ПРАКТИЧЕСКИМ РАБОТАМ 1 И 2"), 16, bold=True)
    set_font(title.add_run("\nПрогнозирование объёма торгов и волатильности BTC/USDT на OKX"), 14)
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(0)
    doc.styles["Title"].font.name = "Times New Roman"
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    # The added subject line replaces one spacer from the sample title page.
    spacer = doc.paragraphs[12]
    if spacer.text:
        raise ValueError("Expected an empty title-page spacer")
    spacer._element.getparent().remove(spacer._element)

    names = doc.tables[1]
    names.rows[0].cells[0].text = "Студент"
    names.rows[0].cells[1].text = "Камалов Т. А." + (f"  {group}" if group else "")
    names.rows[1].cells[1].text = teacher
    for row in names.rows[:2]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, 14)

    title_section = doc.sections[0]
    title_section.top_margin = Cm(2)
    title_section.bottom_margin = Cm(2)
    title_section.left_margin = Cm(3)
    title_section.right_margin = Cm(1)
    for footer in (title_section.footer, title_section.even_page_footer,
                   title_section.first_page_footer):
        for child in list(footer._element):
            footer._element.remove(child)
        footer.add_paragraph()

    set_style(doc, "Normal", 14, first=1.25)
    set_style(doc, "PAS H1", 18, bold=True, first=1.25, after=28.35,
              align=WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc, "PAS H2", 16, bold=True, first=1.25, before=42.55,
              after=28.35, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc, "PAS Special", 18, bold=True, after=28.35,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_style(doc, "PAS TOC", 14, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc, "PAS Caption", 14, italic=True, before=12, after=3,
              align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
    set_style(doc, "PAS Table", 12, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.1)
    set_style(doc, "PAS Source", 14, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)

    toc = add_paragraph(doc, "СОДЕРЖАНИЕ", "PAS Special")
    toc.paragraph_format.page_break_before = True
    for kind, heading in CONTENT:
        if kind not in {"h1", "h2"}:
            continue
        p = add_paragraph(doc, heading + "\t" + str(pages.get(heading, 0)), "PAS TOC")
        if kind == "h2":
            p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(16.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )

    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.footer_distance = Cm(1.25)
    add_page_number(section)
    for kind, value in CONTENT:
        if kind == "h1":
            p = add_paragraph(doc, value, "PAS H1")
            if value.startswith(("1 ПРАКТИЧЕСКАЯ", "2 ПРАКТИЧЕСКАЯ")):
                p.paragraph_format.page_break_before = True
        elif kind == "h2":
            add_paragraph(doc, value, "PAS H2")
        elif kind == "table":
            add_table(doc, value)
        elif kind == "source":
            add_paragraph(doc, value, "PAS Source")
        else:
            add_paragraph(doc, value)
    doc.core_properties.title = "Прогнозирование объёма торгов и волатильности BTC/USDT на OKX"
    doc.core_properties.author = "Камалов Т. А."
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def page_map(pdf_path: Path) -> dict[str, int]:
    reader = PdfReader(pdf_path)
    result = {}
    headings = [value for kind, value in CONTENT if kind in {"h1", "h2"}]
    for number, page in enumerate(reader.pages, start=1):
        if number <= 2:
            continue
        # PDF text extraction may insert spaces inside bold Cyrillic words.
        text = re.sub(r"\s+", "", page.extract_text() or "").upper()
        for heading in headings:
            needle = re.sub(r"\s+", "", heading).upper()
            if heading not in result and needle in text:
                result[heading] = number
    missing = set(headings) - set(result)
    if missing:
        raise RuntimeError(f"TOC headings not found in preliminary PDF: {sorted(missing)}")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--renderer", type=Path, default=DEFAULT_RENDERER)
    parser.add_argument("--group", default="")
    parser.add_argument("--teacher", default="")
    args = parser.parse_args()
    if not args.template.is_file():
        parser.error(f"Title-page template not found: {args.template}")
    with tempfile.TemporaryDirectory(prefix="pas_report_") as temporary:
        temp = Path(temporary)
        draft = temp / "draft.docx"
        render_report(args.template, draft, args.group, args.teacher, {})
        subprocess.run(
            [sys.executable, str(args.renderer), str(draft),
             "--output_dir", str(temp / "render"), "--emit_pdf"],
            check=True, stdout=subprocess.DEVNULL,
        )
        pages = page_map(temp / "render" / "draft.pdf")
        render_report(args.template, args.output, args.group, args.teacher, pages)
    print(args.output)
    print("TOC pages:", pages)


if __name__ == "__main__":
    main()
